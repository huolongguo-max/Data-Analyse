from __future__ import annotations

import sys
from pathlib import Path
from tempfile import TemporaryDirectory

import pandas as pd
from openpyxl import Workbook, load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from reserve_processor.constants import BASE_METRICS, SHEET_DETAIL, SHEET_SCHEDULE, SHEET_TARGET  # noqa: E402
from reserve_processor.excel_io import (  # noqa: E402
    _build_detail_export_frame_with_mode,
    _build_target_export_frame_with_mode,
    write_rounded_export_copy,
    write_xlsx,
)
from reserve_processor.utils import metric_date_col  # noqa: E402


def test_detail_export_keeps_only_base_reserve_daily_columns_and_rounds_to_four_decimals() -> None:
    base_date_col = metric_date_col(BASE_METRICS[0], pd.Timestamp("2026-03-01"))
    high_date_col = metric_date_col(BASE_METRICS[1], pd.Timestamp("2026-03-01"))
    base_revenue_date_col = metric_date_col(BASE_METRICS[2], pd.Timestamp("2026-03-01"))
    high_revenue_date_col = metric_date_col(BASE_METRICS[3], pd.Timestamp("2026-03-01"))
    detail = pd.DataFrame(
        [
            {
                BASE_METRICS[0]: 10,
                base_date_col: 1.23456,
                high_date_col: 2.34567,
                base_revenue_date_col: 9.87654,
                high_revenue_date_col: 8.76543,
            }
        ]
    )

    exported = _build_detail_export_frame_with_mode(detail, round_for_export=False)

    assert exported.iloc[0][BASE_METRICS[0]] == 10
    assert exported.columns.tolist() == [BASE_METRICS[0], "3月1日"]
    assert exported.iloc[0, 1] == 1.2346


def test_write_xlsx_formats_detail_daily_reserve_columns_with_four_decimals() -> None:
    date_col = metric_date_col(BASE_METRICS[0], pd.Timestamp("2026-03-01"))
    detail = pd.DataFrame([{BASE_METRICS[0]: 10, date_col: 1.23456}])
    schedule = pd.DataFrame([{"示例列": 1}])
    target = pd.DataFrame([{"示例列": 1}])

    with TemporaryDirectory() as tmp_dir:
        tmp_path = Path(tmp_dir)
        template_path = tmp_path / "template.xlsx"
        output_path = tmp_path / "output.xlsx"

        workbook = Workbook()
        detail_ws = workbook.active
        detail_ws.title = SHEET_DETAIL
        detail_ws.cell(1, 1, BASE_METRICS[0])
        detail_ws.cell(2, 1, 0)
        workbook.create_sheet(SHEET_SCHEDULE).cell(1, 1, "示例列")
        workbook[SHEET_SCHEDULE].cell(2, 1, 0)
        workbook.create_sheet(SHEET_TARGET).cell(1, 1, "示例列")
        workbook[SHEET_TARGET].cell(2, 1, 0)
        workbook.save(template_path)

        write_xlsx(
            output_path=output_path,
            template_path=template_path,
            detail=detail,
            schedule=schedule,
            target=target,
            round_for_export=False,
        )

        written = load_workbook(output_path)
        sheet = written[SHEET_DETAIL]
        assert sheet.cell(1, 2).value == "3月1日"
        assert sheet.cell(2, 2).value == 1.2346
        assert sheet.cell(2, 2).number_format == "#,##0.0000"


def test_write_rounded_export_copy_normalizes_base_detail_date_header() -> None:
    date_col = metric_date_col(BASE_METRICS[0], pd.Timestamp("2026-03-01"))

    with TemporaryDirectory() as tmp_dir:
        tmp_path = Path(tmp_dir)
        source_path = tmp_path / "source.xlsx"
        output_path = tmp_path / "rounded.xlsx"

        workbook = Workbook()
        detail_ws = workbook.active
        detail_ws.title = SHEET_DETAIL
        detail_ws.cell(1, 1, date_col)
        detail_ws.cell(2, 1, 12.34567)
        workbook.create_sheet(SHEET_SCHEDULE).cell(1, 1, "示例列")
        workbook.create_sheet(SHEET_TARGET).cell(1, 1, "示例列")
        workbook.save(source_path)

        write_rounded_export_copy(
            source_path=source_path,
            template_path=source_path,
            output_path=output_path,
        )

        written = load_workbook(output_path)
        sheet = written[SHEET_DETAIL]
        assert sheet.cell(1, 1).value == "3月1日"
        assert sheet.cell(2, 1).value == 12.3457
        assert sheet.cell(2, 1).number_format == "#,##0.0000"


def test_target_export_renames_entered_column_header() -> None:
    target = pd.DataFrame(
        [
            {
                "目标层面": "片区",
                "目标对象": "华东",
                "预测业绩目标": 100,
                "已进线": 12,
                "已达成预测业绩": 30,
            }
        ]
    )

    exported = _build_target_export_frame_with_mode(target, round_for_export=False)

    assert "已进线量" in exported.columns
    assert "已进线" not in exported.columns
    assert exported.iloc[0]["已进线量"] == 12


def test_write_rounded_export_copy_renames_target_header() -> None:
    with TemporaryDirectory() as tmp_dir:
        tmp_path = Path(tmp_dir)
        source_path = tmp_path / "source.xlsx"
        output_path = tmp_path / "rounded.xlsx"

        workbook = Workbook()
        target_ws = workbook.active
        target_ws.title = SHEET_TARGET
        target_ws.cell(1, 1, "已进线")
        target_ws.cell(2, 1, 12.4)
        workbook.save(source_path)

        write_rounded_export_copy(
            source_path=source_path,
            template_path=source_path,
            output_path=output_path,
        )

        written = load_workbook(output_path)
        sheet = written[SHEET_TARGET]
        assert sheet.cell(1, 1).value == "已进线量"
        assert sheet.cell(2, 1).value == 12
        assert sheet.cell(2, 1).number_format == "#,##0"
