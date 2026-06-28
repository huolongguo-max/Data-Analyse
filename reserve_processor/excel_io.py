# reserve_processor/excel_io.py

from __future__ import annotations

from copy import copy
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
import re

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from .constants import (
    BASE_METRICS,
    RATE_METRICS,
    REQUIRED_SHEETS,
    SHEET_ALIASES,
    SHEET_DETAIL,
    SHEET_SCHEDULE,
    SHEET_TARGET,
    TOTAL_METRICS,
)
from .filename_parser import parse_filename
from .models import WorkbookData
from .utils import is_blank, remove_existing_date_metric_columns, to_number

DETAIL_RATE_HEADERS = {"BD准确率"}
DETAIL_INTEGER_HEADERS = {
    "冲高储备量",
    "周储备提醒",
    "下周储备量",
    "保底储备量",
    "下周保底储备量",
    "冲高预测业绩",
    "保底预测业绩",
}
DETAIL_DAILY_DECIMAL_DIGITS = 4
DETAIL_EXPORT_METRIC = BASE_METRICS[0]
INTERNAL_COLUMN_ALIASES = {
    SHEET_TARGET: {
        "已进线量": "已进线",
    },
}
EXPORT_HEADER_RENAMES = {
    SHEET_TARGET: {
        "已进线": "已进线量",
    },
}
SCHEDULE_INTEGER_HEADERS = {
    "已进线量",
    "已达成预测业绩",
    *BASE_METRICS,
    *TOTAL_METRICS,
}
TARGET_INTEGER_HEADERS = {
    "预测业绩目标",
    "已进线量",
    "已进线",
    "已达成预测业绩",
    *BASE_METRICS,
    *TOTAL_METRICS,
}
DATE_HEADER_PATTERN = re.compile(r"^(?:\d{4}年)?\d{1,2}月\d{1,2}日$")
DETAIL_METRIC_PREFIXES = tuple(f"{metric}-" for metric in BASE_METRICS)


def normalize_sheet_name(sheet_name: str) -> str:
    name = str(sheet_name).strip()
    return SHEET_ALIASES.get(name, name)


def _normalize_internal_sheet_columns(sheet_name: str, frame: pd.DataFrame) -> pd.DataFrame:
    aliases = INTERNAL_COLUMN_ALIASES.get(sheet_name, {})
    rename_map = {
        source: target
        for source, target in aliases.items()
        if source in frame.columns and target not in frame.columns
    }
    if not rename_map:
        return frame
    return frame.rename(columns=rename_map)


def _rename_export_headers(sheet_name: str, frame: pd.DataFrame) -> pd.DataFrame:
    rename_map = EXPORT_HEADER_RENAMES.get(sheet_name, {})
    if not rename_map:
        return frame
    return frame.rename(columns=rename_map)


def read_xlsx(path: str | Path) -> WorkbookData:
    path = Path(path)
    meta = parse_filename(path)

    raw_sheets = pd.read_excel(
        path,
        sheet_name=None,
        dtype=object,
        keep_default_na=False,
        engine="openpyxl",
    )

    sheets = {
        normalize_sheet_name(sheet_name): df
        for sheet_name, df in raw_sheets.items()
    }
    sheets = {
        sheet_name: _normalize_internal_sheet_columns(sheet_name, df)
        for sheet_name, df in sheets.items()
    }

    missing = [sheet for sheet in REQUIRED_SHEETS if sheet not in sheets]
    if missing:
        raise ValueError(
            f"缺少必要 sheet: {missing}；当前 sheet={list(raw_sheets.keys())}"
        )

    return WorkbookData(
        meta=meta,
        detail=sheets[SHEET_DETAIL],
        schedule=sheets[SHEET_SCHEDULE],
        target=sheets[SHEET_TARGET],
    )


def _round_half_up(value, digits: int = 0) -> int | float | None:
    number = to_number(value)
    if number is None:
        return None

    quant = Decimal("1") if digits == 0 else Decimal(f"1e-{digits}")
    rounded = Decimal(str(number)).quantize(quant, rounding=ROUND_HALF_UP)

    if digits == 0:
        return int(rounded)

    return float(rounded)


def _round_rate_ratio_for_export(value) -> float | None:
    number = to_number(value)
    if number is None:
        return None

    pct = (Decimal(str(number)) * Decimal("100")).quantize(
        Decimal("1"),
        rounding=ROUND_HALF_UP,
    )
    return float(pct / Decimal("100"))


def _decimal_number_format(digits: int) -> str:
    return f"#,##0.{''.join('0' for _ in range(digits))}"


def _normalize_export_value(value):
    if is_blank(value):
        return None

    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()

    return value


def _is_export_date_header(header: str) -> bool:
    return bool(DATE_HEADER_PATTERN.match(str(header).strip()))


def _extract_detail_metric_date_part(header: str) -> str | None:
    value = str(header).strip()
    if not value.startswith(DETAIL_METRIC_PREFIXES):
        return None

    _, _, date_part = value.partition("-")
    date_part = date_part.strip()
    if not date_part or not _is_export_date_header(date_part):
        return None

    return date_part


def _is_detail_metric_date_header(header: str) -> bool:
    return _extract_detail_metric_date_part(header) is not None


def _to_detail_export_date_header(header: str) -> str:
    date_part = _extract_detail_metric_date_part(header)
    if date_part is not None:
        return date_part
    return str(header).strip()


def _get_detail_metric_date_columns(detail: pd.DataFrame) -> list[str]:
    return [
        str(col)
        for col in detail.columns
        if _is_detail_metric_date_header(str(col))
        and str(col).strip().startswith(f"{DETAIL_EXPORT_METRIC}-")
    ]


def _build_detail_export_frame(detail: pd.DataFrame) -> pd.DataFrame:
    return _build_detail_export_frame_with_mode(detail, round_for_export=False)


def _build_detail_export_frame_with_mode(detail: pd.DataFrame, *, round_for_export: bool) -> pd.DataFrame:
    base_detail = remove_existing_date_metric_columns(detail).reset_index(drop=True)
    detail_metric_date_cols = _get_detail_metric_date_columns(detail)
    export_dates = detail[detail_metric_date_cols].copy()

    for col in export_dates.columns:
        export_dates[col] = export_dates[col].map(
            lambda value: _round_half_up(value, DETAIL_DAILY_DECIMAL_DIGITS)
        )
    export_dates.columns = [
        _to_detail_export_date_header(col)
        for col in export_dates.columns
    ]

    if round_for_export:
        for col in DETAIL_INTEGER_HEADERS:
            if col in base_detail.columns:
                base_detail[col] = base_detail[col].map(_round_half_up)

        for col in DETAIL_RATE_HEADERS:
            if col in base_detail.columns:
                base_detail[col] = base_detail[col].map(_round_rate_ratio_for_export)

    return pd.concat([base_detail, export_dates], axis=1)


def _build_schedule_export_frame(schedule: pd.DataFrame) -> pd.DataFrame:
    return _build_schedule_export_frame_with_mode(schedule, round_for_export=False)


def _build_schedule_export_frame_with_mode(schedule: pd.DataFrame, *, round_for_export: bool) -> pd.DataFrame:
    out = schedule.copy()

    if round_for_export:
        for col in SCHEDULE_INTEGER_HEADERS:
            if col in out.columns:
                out[col] = out[col].map(_round_half_up)

    return out


def _build_target_export_frame(target: pd.DataFrame) -> pd.DataFrame:
    return _build_target_export_frame_with_mode(target, round_for_export=False)


def _build_target_export_frame_with_mode(target: pd.DataFrame, *, round_for_export: bool) -> pd.DataFrame:
    out = _rename_export_headers(SHEET_TARGET, target.copy())

    if round_for_export:
        for col in TARGET_INTEGER_HEADERS:
            if col in out.columns:
                out[col] = out[col].map(_round_half_up)

        for col in RATE_METRICS:
            if col in out.columns:
                out[col] = out[col].map(_round_rate_ratio_for_export)

    return out


def _copy_cell_style(source_cell, target_cell) -> None:
    target_cell.font = copy(source_cell.font)
    target_cell.fill = copy(source_cell.fill)
    target_cell.border = copy(source_cell.border)
    target_cell.alignment = copy(source_cell.alignment)
    target_cell.protection = copy(source_cell.protection)
    target_cell.number_format = source_cell.number_format


_FONT_HEADER = Font(name="\u5fae\u8f6f\u96c5\u9ed1", size=10, bold=True)   # 微软雅黑 10 加粗
_FONT_DATA = Font(name="\u5b8b\u4f53", size=10, bold=False)                 # 宋体 10 不加粗
_ALIGN_LEFT = Alignment(horizontal="left")


def _apply_export_font(cell, *, is_header: bool) -> None:
    cell.font = _FONT_HEADER if is_header else _FONT_DATA
    cell.alignment = _ALIGN_LEFT


def _build_template_header_map(template_ws) -> dict[str, int]:
    return {
        str(template_ws.cell(1, col_idx).value): col_idx
        for col_idx in range(1, template_ws.max_column + 1)
        if template_ws.cell(1, col_idx).value is not None
    }


def _resolve_style_column(
    *,
    sheet_name: str,
    header: str,
    template_header_map: dict[str, int],
) -> int:
    if header in template_header_map:
        return template_header_map[header]

    for legacy_header, export_header in EXPORT_HEADER_RENAMES.get(sheet_name, {}).items():
        if header == export_header and legacy_header in template_header_map:
            return template_header_map[legacy_header]

    if sheet_name == SHEET_DETAIL:
        return template_header_map.get("保底储备量", 1)

    if sheet_name == SHEET_SCHEDULE:
        if "业绩" in header:
            return template_header_map.get("保底总预测业绩", 1)
        return template_header_map.get("保底总储备量", 1)

    if sheet_name == SHEET_TARGET:
        if "达标率" in header:
            return template_header_map.get("保底总预测业绩达标率", 1)
        if "业绩" in header:
            return template_header_map.get("保底总预测业绩", 1)
        return template_header_map.get("保底总储备量", 1)

    return 1


def _copy_column_width(template_ws, source_col_idx: int, target_ws, target_col_idx: int) -> None:
    source_letter = get_column_letter(source_col_idx)
    target_letter = get_column_letter(target_col_idx)
    source_dim = template_ws.column_dimensions[source_letter]

    if source_dim.width is not None:
        target_ws.column_dimensions[target_letter].width = source_dim.width


def _apply_sheet_defaults(template_ws, target_ws, data_row_count: int) -> None:
    target_ws.freeze_panes = template_ws.freeze_panes
    target_ws.sheet_view.showGridLines = template_ws.sheet_view.showGridLines

    if template_ws.sheet_format.defaultRowHeight is not None:
        target_ws.sheet_format.defaultRowHeight = template_ws.sheet_format.defaultRowHeight

    for row_idx in range(1, data_row_count + 2):
        template_row_idx = row_idx if row_idx <= template_ws.max_row else min(template_ws.max_row, 2)
        if template_row_idx <= 0:
            continue
        source_dim = template_ws.row_dimensions[template_row_idx]
        if source_dim.height is not None:
            target_ws.row_dimensions[row_idx].height = source_dim.height


def _write_sheet_from_frame(
    *,
    workbook,
    sheet_name: str,
    frame: pd.DataFrame,
    template_ws,
    integer_columns: set[str],
    rate_columns: set[str],
    fixed_decimal_columns: dict[str, int],
    round_for_export: bool,
) -> None:
    ws = workbook.create_sheet(sheet_name)
    template_header_map = _build_template_header_map(template_ws)
    style_columns = [
        _resolve_style_column(
            sheet_name=sheet_name,
            header=str(header),
            template_header_map=template_header_map,
        )
        for header in frame.columns
    ]

    _apply_sheet_defaults(template_ws, ws, len(frame))

    for col_idx, header in enumerate(frame.columns, start=1):
        header_cell = ws.cell(1, col_idx, value=header)
        source_col_idx = style_columns[col_idx - 1]
        _copy_cell_style(template_ws.cell(1, source_col_idx), header_cell)
        header_cell.number_format = "General"
        _apply_export_font(header_cell, is_header=True)
        _copy_column_width(template_ws, source_col_idx, ws, col_idx)

    for row_idx, row in enumerate(frame.itertuples(index=False, name=None), start=2):
        template_row_idx = row_idx if row_idx <= template_ws.max_row else min(template_ws.max_row, 2)
        template_row_idx = max(template_row_idx, 1)

        for col_idx, value in enumerate(row, start=1):
            header = str(frame.columns[col_idx - 1])
            source_col_idx = style_columns[col_idx - 1]
            cell = ws.cell(row_idx, col_idx, value=_normalize_export_value(value))
            _copy_cell_style(template_ws.cell(template_row_idx, source_col_idx), cell)
            _apply_export_font(cell, is_header=False)

            if header in fixed_decimal_columns:
                cell.number_format = _decimal_number_format(fixed_decimal_columns[header])
            elif header in integer_columns:
                cell.number_format = "#,##0" if round_for_export else "#,##0.000000"
            elif header in rate_columns:
                cell.number_format = "0%" if round_for_export else "0.000000%"


def write_xlsx(
    output_path: str | Path,
    template_path: str | Path,
    detail: pd.DataFrame,
    schedule: pd.DataFrame,
    target: pd.DataFrame,
    *,
    round_for_export: bool = False,
) -> Path:
    output_path = Path(output_path)
    template_path = Path(template_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    template_wb = load_workbook(template_path)
    export_wb = Workbook()
    export_wb.remove(export_wb.active)

    detail_export = _build_detail_export_frame_with_mode(detail, round_for_export=round_for_export)
    schedule_export = _build_schedule_export_frame_with_mode(schedule, round_for_export=round_for_export)
    target_export = _build_target_export_frame_with_mode(target, round_for_export=round_for_export)

    _write_sheet_from_frame(
        workbook=export_wb,
        sheet_name=SHEET_DETAIL,
        frame=detail_export,
        template_ws=template_wb[SHEET_DETAIL],
        integer_columns=DETAIL_INTEGER_HEADERS,
        rate_columns=DETAIL_RATE_HEADERS,
        fixed_decimal_columns={
            col: DETAIL_DAILY_DECIMAL_DIGITS
            for col in detail_export.columns
            if _is_export_date_header(col)
        },
        round_for_export=round_for_export,
    )
    _write_sheet_from_frame(
        workbook=export_wb,
        sheet_name=SHEET_SCHEDULE,
        frame=schedule_export,
        template_ws=template_wb[SHEET_SCHEDULE],
        integer_columns=SCHEDULE_INTEGER_HEADERS,
        rate_columns=set(),
        fixed_decimal_columns={},
        round_for_export=round_for_export,
    )
    _write_sheet_from_frame(
        workbook=export_wb,
        sheet_name=SHEET_TARGET,
        frame=target_export,
        template_ws=template_wb[SHEET_TARGET],
        integer_columns=TARGET_INTEGER_HEADERS,
        rate_columns=set(RATE_METRICS),
        fixed_decimal_columns={},
        round_for_export=round_for_export,
    )

    export_wb.save(output_path)

    return output_path


def write_rounded_export_copy(
    source_path: str | Path,
    template_path: str | Path,
    output_path: str | Path,
) -> Path:
    source_path = Path(source_path)
    template_path = Path(template_path)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = load_workbook(source_path)
    for sheet_name in (SHEET_DETAIL, SHEET_SCHEDULE, SHEET_TARGET):
        if sheet_name not in workbook.sheetnames:
            continue

        ws = workbook[sheet_name]

        # apply font and alignment to every cell in the sheet
        for col_idx in range(1, ws.max_column + 1):
            _apply_export_font(ws.cell(1, col_idx), is_header=True)
            for row_idx in range(2, ws.max_row + 1):
                _apply_export_font(ws.cell(row_idx, col_idx), is_header=False)

        # round numeric columns
        for col_idx in range(1, ws.max_column + 1):
            header_value = ws.cell(1, col_idx).value
            if header_value is None:
                continue

            header = str(header_value)
            if (
                sheet_name == SHEET_DETAIL
                and header.strip().startswith(f"{DETAIL_EXPORT_METRIC}-")
            ):
                normalized_detail_header = _to_detail_export_date_header(header)
                if normalized_detail_header != header:
                    ws.cell(1, col_idx).value = normalized_detail_header
                    header = normalized_detail_header
            renamed_header = EXPORT_HEADER_RENAMES.get(sheet_name, {}).get(header, header)
            if renamed_header != header:
                ws.cell(1, col_idx).value = renamed_header
            header = renamed_header
            is_detail_daily_col = (
                sheet_name == SHEET_DETAIL and _is_export_date_header(header)
            )
            is_rate_col = (
                (sheet_name == SHEET_DETAIL and header in DETAIL_RATE_HEADERS)
                or (sheet_name == SHEET_TARGET and header in RATE_METRICS)
            )
            is_amount_col = (
                (sheet_name == SHEET_DETAIL and (header in DETAIL_INTEGER_HEADERS or is_detail_daily_col))
                or (sheet_name == SHEET_SCHEDULE and header in SCHEDULE_INTEGER_HEADERS)
                or (sheet_name == SHEET_TARGET and header in TARGET_INTEGER_HEADERS)
            )

            if not is_rate_col and not is_amount_col:
                continue

            for row_idx in range(2, ws.max_row + 1):
                cell = ws.cell(row_idx, col_idx)
                if is_rate_col:
                    cell.value = _round_rate_ratio_for_export(cell.value)
                    cell.number_format = "0%"
                elif is_detail_daily_col:
                    cell.value = _round_half_up(cell.value, DETAIL_DAILY_DECIMAL_DIGITS)
                    cell.number_format = _decimal_number_format(DETAIL_DAILY_DECIMAL_DIGITS)
                else:
                    cell.value = _round_half_up(cell.value)
                    cell.number_format = "#,##0"

    workbook.save(output_path)
    return output_path
